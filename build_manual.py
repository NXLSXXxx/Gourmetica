from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\xampp\htdocs\Gourmetica")
IMG = ROOT / "storage" / "app"
OUT = ROOT / "Manual_Usuario_Gourmetica.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Página ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)

def add_image_section(doc, title, route, image_name, instructions, notes=None):
    doc.add_heading(title, level=1)
    p = doc.add_paragraph()
    p.add_run('Ruta: ').bold = True
    p.add_run(route)
    doc.add_paragraph(instructions)
    image_path = IMG / image_name
    if image_path.exists():
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width = Inches(5.15) if image_name.startswith('manual4-') or image_name == 'manual-dashboard-wide.png' else Inches(6.35)
        pic.add_run().add_picture(str(image_path), width=width)
        cap = doc.add_paragraph(f'Figura. Pantalla de {title.lower()}.')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph('[Captura no disponible en el entorno de demostración.]')
    if notes:
        p = doc.add_paragraph()
        p.add_run('Consideraciones: ').bold = True
        p.add_run(notes)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
for name, size, color in [('Title', 26, '000000'), ('Heading 1', 17, '000000'), ('Heading 2', 13, '000000')]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name != 'Title' else 'Aptos Display'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)

header = sec.header.paragraphs[0]
header.text = 'GOURMETICA  |  MANUAL DE USUARIO'
header.runs[0].font.size = Pt(8)
header.runs[0].font.bold = True
header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
footer = sec.footer.paragraphs[0]
add_page_number(footer)
footer.runs[0].font.size = Pt(8)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.space_after = Pt(24)
p.add_run('GOURMETICA')
p.runs[0].font.name = 'Georgia'
p.runs[0].font.size = Pt(28)
p.runs[0].font.bold = True
p.runs[0].font.color.rgb = RGBColor(183, 122, 62)
title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Manual de Usuario Gourmetica')
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('Guía visual de la tienda en línea y la intranet administrativa').italic = True
doc.add_paragraph()
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.add_run('Versión documentada: octubre de 2026\n').bold = True
intro.add_run('Este manual explica los recorridos principales del sistema Gourmetica mediante capturas de pantalla reales y pasos de uso para clientes, caja, operaciones y administración.')
doc.add_page_break()

doc.add_heading('Contenido del sistema', level=1)
toc = [
    ('1', 'Acceso público y compras'),
    ('2', 'Intranet administrativa'),
    ('3', 'Incidencia conocida'),
    ('4', 'Buenas prácticas de operación'),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['Sección', 'Contenido']):
    cell.text = text; shade(cell, '374151'); set_cell_border(cell); set_cell_margins(cell)
    for r in cell.paragraphs[0].runs:
        r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255)
for i, item in enumerate(toc):
    cells = table.add_row().cells
    cells[0].text, cells[1].text = item
    for c in cells:
        set_cell_border(c); set_cell_margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i % 2: shade(c, 'F3F4F6')
doc.add_paragraph('Roles principales: cliente, cajero, administrador de sede, administrador general e ingeniero de sistemas. Las opciones visibles pueden variar según el rol y los permisos asignados.')
doc.add_page_break()

doc.add_heading('1 Acceso público y compras', level=1)
doc.add_paragraph('La parte pública permite consultar la propuesta de Gourmetica, seleccionar una sede, revisar el catálogo, agregar productos al carrito, solicitar catering y consultar información de contacto.')

public = [
    ('Página de inicio', '/', 'manual-home.png', 'La página de inicio presenta la marca, la propuesta de valor, las categorías y el acceso rápido a la carta. Use “Ver Carta” o “Hacer Pedido” para comenzar una compra.', 'La sede seleccionada influye en disponibilidad y precios.'),
    ('Catálogo de productos', '/shop', 'manual-catalogo.png', 'Explore las categorías y use la búsqueda para encontrar productos. Seleccione una tarjeta para revisar el detalle y agregar el producto al carrito.', 'El catálogo puede mostrar productos diferentes según la sede elegida.'),
    ('Carrito de compras', '/cart', 'manual-carrito.png', 'Revise cantidades, precios y productos seleccionados. Desde aquí puede retirar artículos o continuar hacia el checkout.', 'Verifique la sede y el contenido antes de continuar.'),
    ('Nuestras sedes', '/locations', 'manual-sedes.png', 'Consulte las direcciones, horarios y datos de cada sede. Seleccione la sede que corresponda al recojo o pedido.', 'La selección de sede es importante para confirmar disponibilidad.'),
    ('Catering y eventos', '/catering', 'manual-catering.png', 'Use el formulario para solicitar una propuesta de catering. Indique datos de contacto, fecha, cantidad de personas y necesidades del evento.', 'La solicitud queda disponible para seguimiento administrativo.'),
    ('Contacto', '/contact', 'manual-contacto.png', 'Consulte los canales de atención y comuníquese con Gourmetica para dudas, pedidos personalizados o soporte.', 'Los datos mostrados corresponden al entorno documentado.'),
    ('Inicio de sesión de cliente', '/login', 'manual-login.png', 'Ingrese su correo y contraseña para acceder a su perfil, consultar pedidos y utilizar funciones asociadas a la cuenta.', 'El registro de cliente se realiza desde la opción de creación de cuenta.'),
]
for data in public: add_image_section(doc, *data)

doc.add_heading('2 Intranet administrativa', level=1)
doc.add_paragraph('La intranet concentra la operación interna. El menú lateral se organiza por ventas y clientes, inventario y cocina, logística y catálogo, y configuración del sistema. El acceso depende del rol del usuario.')
add_image_section(doc, 'Acceso administrativo', '/admin/login', 'manual-admin-login.png', 'Ingrese con una cuenta administrativa autorizada. Después de validar las credenciales, el sistema redirige al panel de control.', 'No comparta las credenciales de producción. Use cuentas individuales y cierre sesión al terminar.')

admin = [
    ('Panel de control', '/intranet/dashboard', 'manual-dashboard-wide.png', 'Revise indicadores de ventas, stock, compras, sedes activas y el resumen de flujo de caja. Use los accesos del menú para profundizar en cada área.', 'Los indicadores dependen de los datos registrados y del alcance del usuario.'),
    ('Pedidos en vivo y caja', '/intranet/live-orders', 'manual4-live-orders.png', 'Gestione ventas rápidas y pedidos activos desde caja. Registre la venta, revise el detalle y actualice el estado conforme avanza la atención.', 'Confirme el pago y la sede antes de cerrar la operación.'),
    ('Pedidos', '/intranet/orders', 'manual4-pedidos.png', 'Consulte pedidos del canal en línea, abra el detalle y actualice su estado. Desde el detalle se pueden revisar productos, cliente y entrega.', 'Los cambios de estado deben reflejar el avance real del pedido.'),
    ('Solicitudes de catering', '/intranet/catering', 'manual4-catering-admin.png', 'Revise solicitudes recibidas, abra el detalle y cambie su estado para coordinar la atención comercial y operativa.', 'Mantenga actualizados los datos de contacto y la fecha del evento.'),
    ('Ventas y SUNAT', '/intranet/sales', 'manual4-ventas.png', 'Consulte ventas registradas y revise el estado de declaración. Use las acciones disponibles según el permiso asignado.', 'La declaración debe realizarse solo después de validar el comprobante.'),
    ('Clientes', '/intranet/clients', 'manual4-clientes.png', 'Busque clientes y consulte su historial de pedidos. Esta pantalla sirve para atención y seguimiento comercial.', 'Trate los datos de clientes como información reservada.'),
    ('Insumos', '/intranet/supplies', 'manual4-insumos.png', 'Registre y consulte materias primas, unidades y stock. Use el ajuste de stock para actualizar existencias cuando corresponda.', 'Registre movimientos con una razón clara para facilitar la auditoría.'),
    ('Producción y cocina', '/intranet/productions', 'manual4-produccion.png', 'Consulte preparaciones registradas y cree nuevas producciones según las recetas y disponibilidad de insumos.', 'Revise cantidades antes de registrar una producción.'),
    ('Compras', '/intranet/purchases', 'manual4-compras.png', 'Registre compras a proveedores y consulte los detalles para mantener control de egresos e inventario.', 'Conserve los datos del proveedor y el comprobante de respaldo.'),
    ('Productos', '/intranet/products', 'manual4-productos.png', 'Administre los productos del catálogo: nombre, categoría, descripción, precio, imagen, estado y disponibilidad por sede.', 'Una modificación aquí puede impactar directamente la tienda pública.'),
    ('Categorías', '/intranet/categories', 'manual4-categorias.png', 'Cree, edite y ordene categorías para organizar el catálogo de productos.', 'Use nombres consistentes y fáciles de reconocer por el cliente.'),
    ('Sedes', '/intranet/headquarters', 'manual4-sedes-admin.png', 'Administre sedes, direcciones, contactos y estado de atención. Desde el detalle puede revisar la relación de productos con cada sede.', 'Verifique la información antes de activar una sede.'),
    ('Zonas de delivery', '/intranet/delivery-zones', 'manual4-zonas-delivery.png', 'Configure zonas, tarifas y condiciones de delivery para calcular correctamente el costo de entrega.', 'Mantenga actualizadas las zonas y tarifas vigentes.'),
    ('Banners', '/intranet/banners', 'manual4-banners.png', 'Administre imágenes y mensajes promocionales visibles en la página principal.', 'Use imágenes optimizadas y textos breves para mantener una buena experiencia.'),
    ('Administradores', '/intranet/users', 'manual4-usuarios.png', 'Cree y edite cuentas internas, asigne roles y, cuando corresponda, vincule una sede.', 'Asigne el menor nivel de permiso necesario para cada función.'),
    ('Configuración', '/intranet/settings', 'manual4-configuracion.png', 'Actualice parámetros generales del sistema, como datos de contacto y configuraciones operativas habilitadas.', 'Cambios de configuración deben revisarse antes de publicarse.'),
]
for data in admin: add_image_section(doc, *data)

doc.add_heading('3 Incidencia conocida', level=1)
doc.add_paragraph('Durante la revisión visual, la ruta /checkout presentó un error de servidor porque la vista resources/views/shop/checkout.blade.php utiliza la variable $selectedHq sin recibirla desde el controlador. Por ese motivo, la pantalla de checkout no se incluye como flujo operativo validado en esta versión del manual.')
p = doc.add_paragraph()
p.add_run('Acción recomendada: ').bold = True
p.add_run('definir o pasar $selectedHq en ShopController@checkout y volver a probar el flujo de selección de sede, cálculo de delivery y pago antes de usar checkout en producción.')

doc.add_heading('4 Buenas prácticas de operación', level=1)
for text in [
    'Use cuentas personales y cierre sesión cuando termine.',
    'Seleccione la sede correcta antes de registrar ventas o revisar disponibilidad.',
    'Revise los datos de una operación antes de confirmar cambios de estado.',
    'Documente ajustes de stock y compras con información suficiente para auditoría.',
    'Antes de publicar productos, categorías o banners, compruebe cómo se verán en la tienda pública.',
    'Realice una prueba controlada de checkout después de corregir la incidencia indicada.'
]:
    doc.add_paragraph(text, style='List Bullet')

doc.add_paragraph('Fin del manual.')
doc.core_properties.title = 'Manual de Usuario Gourmetica'
doc.core_properties.subject = 'Guía visual del sistema Gourmetica'
doc.core_properties.author = 'Gourmetica'
doc.save(OUT)
print(OUT)
